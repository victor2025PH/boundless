# -*- coding: utf-8 -*-
"""翻译媒体链五合一实弹验证（运维显式跑；2026-08-18 P1-P4 装载验证时定型）。

覆盖：语音上传转写翻译（真人参考音+逐字稿重叠断言）/ SRT 保时间轴 / PPTX 保版式 /
视频免上传抽轨转写 / 图片贴回（真 VLM）。全部经页面内同源 fetch（CSRF 安全）。

⚠ **实弹消耗**（176 ASR + VLM + 翻译引擎各一发）——**刻意不进 gate_sweep**，
重启装载翻译链改动后由骑手手动跑一次：

    python tools/live_verify_xlate_chains.py

素材=生产参考音（voice_refs/chen_meiling）、合成 2-cue SRT、内存 PPTX、
参考音 mux 的 mp4、tmp_bbox_spike/src.png（缺失时先跑 tmp_bbox_spike/spike.py 生成）。
预期 12/12 PASS（SRT 时间戳档在 176 部署 verbose_json 前=诚实缺席也算 PASS）。
"""
import base64
import json
import subprocess
import sys
import time
from io import BytesIO
from pathlib import Path

ENGINE = Path(__file__).resolve().parent
sys.path.insert(0, str(ENGINE / "tools"))
from verify_inbox_density import DEFAULT_BASE, DEFAULT_DATA_ROOT, read_token  # noqa: E402

from playwright.sync_api import sync_playwright  # noqa: E402

WAV_RAW = Path(r"D:\chengjie-instances\zhiliao\data\config\voice_refs\chen_meiling.wav")
# 上传口经 16k 单声道版（b64 <2MB）：全局 body 上限 2MB 的代码级放宽随 22:30 窗生效，
# 空窗期用小文件验证链路本身；上限放宽由 test_intent_tags_rate_limit 新用例钉行为。
WAV = Path(__file__).parent / "tmp_bbox_spike" / "probe_16k.wav"
SIDECAR = Path(r"D:\chengjie-instances\zhiliao\data\config\voice_refs\chen_meiling.txt")
SRC_PNG = ENGINE / "tmp_bbox_spike" / "src.png"

results = []


def chk(ok, name, detail=""):
    results.append(ok)
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}  {detail}")


def post_json(page, url, payload):
    return page.evaluate(
        """async ({url, payload}) => {
          const r = await fetch(url, {method:'POST',
            headers:{'Content-Type':'application/json'},
            body: JSON.stringify(payload)});
          try { return await r.json(); } catch(e) { return {http:r.status}; }
        }""", {"url": url, "payload": payload})


def fetch_bytes(page, url):
    b64 = page.evaluate(
        """async (url) => {
          const r = await fetch(url);
          const buf = await r.arrayBuffer();
          let s=''; const u8=new Uint8Array(buf);
          for (let i=0;i<u8.length;i+=0x8000)
            s += String.fromCharCode.apply(null, u8.subarray(i, i+0x8000));
          return btoa(s);
        }""", url)
    return base64.b64decode(b64)


def main() -> int:
    token = read_token(DEFAULT_DATA_ROOT)
    sidecar_text = SIDECAR.read_text(encoding="utf-8").strip()
    print(f"参考音逐字稿: {sidecar_text[:50]}…" if len(sidecar_text) > 50
          else f"参考音逐字稿: {sidecar_text}")

    # 上传口用 16k 单声道小文件（b64 <2MB，见 WAV 注释）
    subprocess.run(["ffmpeg", "-y", "-i", str(WAV_RAW), "-ac", "1", "-ar", "16000",
                    str(WAV)], capture_output=True, timeout=60)
    # 合成一个带真人语音的 mp4（黑屏视频轨 + **原始**参考音音轨——视频链自己抽轨转 16k）
    mp4 = ENGINE / "tmp_bbox_spike" / "verify_speech.mp4"
    r = subprocess.run(
        ["ffmpeg", "-y", "-i", str(WAV_RAW), "-f", "lavfi",
         "-i", "color=c=black:s=64x64:r=5", "-shortest",
         "-pix_fmt", "yuv420p", "-c:a", "aac", str(mp4)],
        capture_output=True, timeout=120)
    if r.returncode != 0:
        print("[ABORT] mp4 合成失败", (r.stderr or b"")[-200:])
        return 1

    srt_src = ("1\n00:00:01,000 --> 00:00:03,000\nHello, how are you today?\n\n"
               "2\n00:00:04,500 --> 00:00:06,200\nPlease send me the invoice.\n\n")

    import pptx as pptx_lib
    from pptx.util import Inches
    prs = pptx_lib.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Quarterly sales report"
    pbuf = BytesIO()
    prs.save(pbuf)

    with sync_playwright() as pw:
        b = pw.chromium.launch()
        ctx = b.new_context()
        ctx.request.post(DEFAULT_BASE + "/login", form={"auth_token": token})
        page = ctx.new_page()
        page.set_default_timeout(300000)
        page.goto(DEFAULT_BASE + "/workspace", wait_until="domcontentloaded")

        # ── 1. 语音上传（真人语音；也替本进程赢下 audio_pipeline 单例竞态）──
        t0 = time.time()
        d = post_json(page, "/api/unified-inbox/translate-voice", {
            "audio_b64": "data:audio/wav;base64,"
                         + base64.b64encode(WAV.read_bytes()).decode(),
            "target_lang": "en", "srt": True})
        dt = time.time() - t0
        tr = str(d.get("transcript") or "")
        dst = str(((d.get("translation") or {}).get("translated_text")) or "")
        overlap = sum(1 for ch in set(tr) if ch in sidecar_text)
        chk(bool(d.get("ok")) and len(tr) >= 8,
            "语音上传：真人语音转写", f"{dt:.1f}s transcript={tr[:32]!r}")
        chk(overlap >= max(4, len(set(tr)) // 3),
            "转写与参考音逐字稿实质重叠", f"overlap_chars={overlap}")
        chk(bool(dst) and dst != tr, "转写文本已译为英文", f"dst={dst[:48]!r}")
        # 诚实缺席=绝无伪造 srt_text；srt_reason 显式化随 22:30 窗装载（本进程还是老路由）
        chk("srt_text" not in d and d.get("srt_reason") in (None, "no_segments"),
            "无分段时诚实缺席字幕（176 未部署 verbose_json）",
            f"srt_reason={d.get('srt_reason')!r}")

        # ── 2. SRT 文档保时间轴 ─────────────────────────────────────
        d = post_json(page, "/api/unified-inbox/translate-document-file", {
            "file_b64": "data:application/x-subrip;base64,"
                        + base64.b64encode(srt_src.encode()).decode(),
            "filename": "verify.srt", "target_lang": "zh", "stream": False})
        ok_srt = bool(d.get("ok")) and d.get("kind") == "file"
        chk(ok_srt, "SRT 上传翻译（stream=false）",
            f"file={d.get('filename')!r} stats={d.get('stats')}")
        if ok_srt:
            data = fetch_bytes(page, d["download_url"]).decode("utf-8")
            chk("00:00:01,000 --> 00:00:03,000" in data
                and "00:00:04,500 --> 00:00:06,200" in data,
                "时间轴逐字节保留")
            chk("Hello, how are you today?" not in data
                and any("\u4e00" <= c <= "\u9fff" for c in data),
                "文本行已译为中文", data.replace("\n", "/")[:120])

        # ── 3. PPTX 保版式 ──────────────────────────────────────────
        d = post_json(page, "/api/unified-inbox/translate-document-file", {
            "file_b64": "data:application/octet-stream;base64,"
                        + base64.b64encode(pbuf.getvalue()).decode(),
            "filename": "verify.pptx", "target_lang": "zh", "stream": False})
        ok_pptx = bool(d.get("ok")) and d.get("kind") == "file"
        chk(ok_pptx, "PPTX 上传翻译", f"stats={d.get('stats')}")
        if ok_pptx:
            data = fetch_bytes(page, d["download_url"])
            reopened = pptx_lib.Presentation(BytesIO(data))
            texts = [p.text for s in reopened.slides for sh in s.shapes
                     if sh.has_text_frame for p in sh.text_frame.paragraphs]
            chk(any(t and t != "Quarterly sales report" for t in texts),
                "PPTX 文本已译且可重新打开", f"texts={texts}")

        # ── 4. 视频（免上传路径：media_ref 直指本地 mp4）────────────
        t0 = time.time()
        d = post_json(page, "/api/unified-inbox/translate-message-media", {
            "conversation_id": "", "message_id": "",
            "media_ref": str(mp4), "media_type": "video", "target_lang": "en"})
        dt = time.time() - t0
        tr_v = str(d.get("transcript") or "")
        chk(bool(d.get("ok")) and d.get("media_kind") == "video" and len(tr_v) >= 8,
            "视频消息转写翻译（ffmpeg 抽轨→ASR→译）",
            f"{dt:.1f}s dur={d.get('video_duration_sec')} transcript={tr_v[:32]!r}")

        # ── 5. 图片贴回（真 VLM + 真翻译 + 回绘）────────────────────
        t0 = time.time()
        d = post_json(page, "/api/unified-inbox/translate-message-media", {
            "conversation_id": "", "message_id": "",
            "media_ref": str(SRC_PNG), "media_type": "image",
            "patch": True, "target_lang": "zh"})
        dt = time.time() - t0
        pb64 = str(d.get("patched_image_b64") or "")
        ok_patch = bool(d.get("ok")) and pb64.startswith("data:image/png;base64,")
        chk(ok_patch, "图片贴回（patch=true 全链）",
            f"{dt:.1f}s stats={d.get('stats')}")
        if ok_patch:
            png = base64.b64decode(pb64.split(",", 1)[1])
            chk(png[:4] == b"\x89PNG", "贴回产物 magic bytes=PNG")
            out = ENGINE / "tmp_bbox_spike" / "route_patched.png"
            out.write_bytes(png)
            print(f"  贴回样张: {out}")

        b.close()
    n_ok = sum(1 for x in results if x)
    print(f"== 装载后实弹验证: {n_ok}/{len(results)} PASS ==")
    return 0 if n_ok == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
