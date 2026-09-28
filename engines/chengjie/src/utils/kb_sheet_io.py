# -*- coding: utf-8 -*-
"""知识库表格导入 / 导入模板（CSV · XLSX · JSON）。

- 表头宽容：中英文列名都认，忽略空格、星号和括号里的提示（「标题（必填）」≡ title）；
  「问题 / 答案」两列的极简 FAQ 表也能直接导。
- 单元格留空 = 不提供该字段：更新模式下不会把已有内容清空（旧实现把空串照写回库）。
- 行级问题带 Excel 行号与原因码回传（``no_title`` / ``bad_reply_mode`` / ``bad_enabled``），
  文案由前端按原因码取 i18n。
- 模板示例与抽屉「新建条目」模板 chip 同源（``new_entry_templates``），不维护第二份示例。
"""
from __future__ import annotations

import csv
import io
import json
import re
from typing import Dict, List, Optional, Tuple

EXAMPLE_TITLE_PREFIXES = ("示例：", "示例:", "【示例】", "Example:", "example:")
TEMPLATE_EXAMPLE_PREFIX = "示例："

# (字段, 中文列名, 必填, 填写说明) —— 模板列顺序即此顺序
SHEET_COLUMNS: List[Tuple[str, str, bool, str]] = [
    ("title", "标题", True, "必填。一条知识的名字，同名视为重复（按导入时选的策略跳过或更新）"),
    ("category", "分类", False, "建议用知识库页已有的分类名；留空归入「其他」"),
    ("triggers", "触发词", False, "客户可能的问法或关键词，多个用逗号、分号或顿号分隔"),
    ("scenario", "使用场景", False, "什么情况下用这条知识，帮 AI 判断要不要引用"),
    ("steps", "处理步骤", False, "处理这类问题的步骤"),
    ("principles", "回复原则", False, "回复时要遵守的要点"),
    ("example_reply_zh", "标准回复", False, "推荐的回复写法；多条示例之间用单独一行 --- 分隔"),
    ("forbidden", "禁止说的话", False, "回复中绝不能出现的说法"),
    ("reply_mode", "回复模式", False, "AI参考（默认）/ AI严格 / 直接输出"),
    ("negative_triggers", "排除词", False, "消息里含这些词时不命中本条，分隔方式同触发词"),
    ("enabled", "启用", False, "是 / 否；留空 = 新条目启用、已有条目不变"),
]

_HEADER_ALIASES: Dict[str, str] = {}
for _field, _zh, _req, _hint in SHEET_COLUMNS:
    _HEADER_ALIASES[_field] = _field
    _HEADER_ALIASES[_zh] = _field
for _alias, _field in {
    "问题": "title", "问法": "title", "名称": "title", "question": "title",
    "类别": "category",
    "trigger": "triggers", "关键词": "triggers", "keywords": "triggers", "keyword": "triggers",
    "场景": "scenario", "适用场景": "scenario",
    "步骤": "steps",
    "原则": "principles",
    "答案": "example_reply_zh", "回答": "example_reply_zh", "回复": "example_reply_zh",
    "示例回复": "example_reply_zh", "answer": "example_reply_zh", "reply": "example_reply_zh",
    "禁忌": "forbidden", "禁用语": "forbidden",
    "模式": "reply_mode",
    "否定触发词": "negative_triggers",
    "是否启用": "enabled", "状态": "enabled",
    "template_key": "template_key", "fallback_group": "fallback_group",
    "reply_direct_spec": "reply_direct_spec",
}.items():
    _HEADER_ALIASES[_alias] = _field

# 标题列是「问题」类表头 = FAQ 表：问题本身就是客户问法，触发词留空时拿它兜底
# （无触发词的条目检索不到，见 kb_store 健康检查 no_triggers）
_QUESTION_HEADERS = {"问题", "问法", "question"}

_REPLY_MODES = {
    "ai_guided": "ai_guided", "ai参考": "ai_guided", "参考": "ai_guided", "ai引导": "ai_guided",
    "ai_strict": "ai_strict", "ai严格": "ai_strict", "严格": "ai_strict",
    "direct": "direct", "直接输出": "direct", "直接": "direct",
}
_REPLY_MODE_LABELS = {"ai_guided": "AI参考", "ai_strict": "AI严格", "direct": "直接输出"}

_TRUE = {"1", "是", "y", "yes", "true", "启用", "开", "on", "√", "✓"}
_FALSE = {"0", "否", "n", "no", "false", "停用", "关", "off", "×", "✗"}

_TERM_SPLIT = re.compile(r"[;；,，、|\n]+")


def split_terms(value) -> List[str]:
    """触发词 / 排除词归一成列表：接受 list、JSON 数组字符串或用 , ， ; ； 、 | 换行分隔的文本。"""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(t).strip() for t in value if str(t).strip()]
    s = str(value).strip()
    if not s:
        return []
    if s.startswith("["):
        try:
            parsed = json.loads(s)
        except ValueError:
            parsed = None   # 不是 JSON 数组，比如 "[VIP] 退款"，按普通分隔符切
        if isinstance(parsed, list):
            return [str(t).strip() for t in parsed if str(t).strip()]
    return [t.strip() for t in _TERM_SPLIT.split(s) if t.strip()]


def normalize_header(h) -> str:
    s = str(h or "").strip().lstrip("\ufeff")
    s = re.sub(r"[（(][^）)]*[）)]", "", s)
    s = re.sub(r"[\s*＊]", "", s)
    return s.lower()


def is_example_title(title: str) -> bool:
    return str(title or "").strip().startswith(EXAMPLE_TITLE_PREFIXES)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, bool):
        return "是" if v else "否"
    return str(v).strip()


def read_csv_table(text: str) -> Tuple[List[str], List[List[str]]]:
    text = (text or "").lstrip("\ufeff")
    first = text.split("\n", 1)[0]
    delim = "\t" if ("\t" in first and "," not in first) else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if not rows:
        return [], []
    return [_cell(h) for h in rows[0]], [[_cell(c) for c in r] for r in rows[1:]]


def read_xlsx_table(data: bytes) -> Tuple[List[str], List[List[str]]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb["知识条目"] if "知识条目" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if header is None:
            return [], []
        return [_cell(h) for h in header], [[_cell(c) for c in r] for r in it]
    finally:
        wb.close()


def rows_to_entries(header: List[str], rows: List[List[str]], *,
                    skip_examples: bool = True) -> Tuple[List[Dict], Dict]:
    """表格 → 条目列表 + 解析报告。每个条目带 ``_row``（Excel 行号，表头为第 1 行）。

    报告：``{columns: {列名: 字段}, unknown_headers, missing_title_column,
    examples_skipped, errors: [{row, title, reason}], warnings: [...]}``；
    ``errors`` 是整行未导入的，``warnings`` 是该行导入了但某个单元格被忽略。
    """
    colmap: Dict[int, str] = {}
    columns: Dict[str, str] = {}
    unknown: List[str] = []
    for i, h in enumerate(header):
        if not str(h).strip():
            continue
        field = _HEADER_ALIASES.get(normalize_header(h))
        if field and field not in colmap.values():
            colmap[i] = field
            columns[str(h)] = field
        else:
            unknown.append(str(h))
    title_header = next((normalize_header(header[i]) for i, f in colmap.items() if f == "title"), "")
    question_style = title_header in _QUESTION_HEADERS
    report: Dict = {
        "columns": columns, "unknown_headers": unknown,
        "missing_title_column": "title" not in colmap.values(),
        "examples_skipped": 0, "errors": [], "warnings": [],
    }
    entries: List[Dict] = []
    if report["missing_title_column"]:
        return entries, report
    for ri, raw in enumerate(rows, start=2):
        cells = {f: (raw[i] if i < len(raw) else "") for i, f in colmap.items()}
        if not any(cells.values()):
            continue
        title = cells.get("title", "")
        if not title:
            report["errors"].append({"row": ri, "title": "", "reason": "no_title"})
            continue
        if skip_examples and is_example_title(title):
            report["examples_skipped"] += 1
            continue
        entry: Dict = {"_row": ri}
        for field, val in cells.items():
            if not val:
                continue
            if field in ("triggers", "negative_triggers"):
                entry[field] = split_terms(val)
            elif field == "reply_mode":
                mode = _REPLY_MODES.get(re.sub(r"\s", "", val).lower())
                if mode:
                    entry[field] = mode
                else:
                    report["warnings"].append({"row": ri, "title": title, "reason": "bad_reply_mode", "value": val})
            elif field == "enabled":
                v = val.strip().lower()
                if v in _TRUE:
                    entry["enabled"] = 1
                elif v in _FALSE:
                    entry["enabled"] = 0
                else:
                    report["warnings"].append({"row": ri, "title": title, "reason": "bad_enabled", "value": val})
            else:
                entry[field] = val
        if not entry.get("triggers"):
            if question_style:
                entry["triggers"] = [title]
                entry["_auto_triggers"] = True
            else:
                report["warnings"].append({"row": ri, "title": title, "reason": "no_triggers"})
        entries.append(entry)
    return entries, report


# ── 模板 ─────────────────────────────────────────────────────

def template_examples(business_domain: Optional[str] = None,
                      categories: Optional[List[str]] = None, limit: int = 2,
                      kinds: Optional[List[str]] = None) -> List[Dict]:
    """导入模板里的示例行。``kinds``（P1-2）含 support 时前置客服示例（活动 / 注册），
    让做客服的客户下载模板就看到自己那类问题该怎么填。"""
    from src.utils.kb_store import new_entry_templates
    out = []
    for t in new_entry_templates(business_domain, categories, kinds=kinds)[:limit]:
        t = dict(t)
        t["title"] = TEMPLATE_EXAMPLE_PREFIX + str(t.get("title", "")).strip()
        out.append(t)
    return out


def _template_row(ex: Dict) -> List[str]:
    row = []
    for field, _zh, _req, _hint in SHEET_COLUMNS:
        v = ex.get(field)
        if field in ("triggers", "negative_triggers"):
            row.append("，".join(split_terms(v)))
        elif field == "reply_mode":
            row.append(_REPLY_MODE_LABELS.get(str(v or "ai_guided"), "AI参考"))
        elif field == "enabled":
            row.append("是")
        else:
            row.append("" if v is None else str(v))
    return row


def build_csv_template(examples: List[Dict]) -> str:
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow([zh for _f, zh, _r, _h in SHEET_COLUMNS])
    for ex in examples:
        w.writerow(_template_row(ex))
    return "\ufeff" + out.getvalue()


def build_json_template(examples: List[Dict]) -> Dict:
    fields = [f for f, _zh, _r, _h in SHEET_COLUMNS]
    entries = []
    for ex in examples:
        e = {f: ex.get(f) for f in fields if ex.get(f) not in (None, "", [])}
        e["triggers"] = split_terms(ex.get("triggers"))
        e.setdefault("reply_mode", "ai_guided")
        e["enabled"] = 1
        entries.append(e)
    return {"version": "1.0", "entries": entries}


TEMPLATE_NOTES = [
    "以「示例：」开头的行导入时默认跳过，可以照着改写，也可以直接删掉。",
    "第一行是列名，请不要修改；不需要的列可以整列删除，只有「标题」必须保留。",
    "单元格留空 = 不填这一项；选「更新已有」导入时，留空的列不会清掉已有内容。",
    "同名（标题完全相同）的条目视为重复，导入时可选择跳过或更新。",
    "直接上传本 .xlsx 文件即可；若另存为 CSV，请选「CSV UTF-8（逗号分隔）」。",
    "已有 FAQ 表只有「问题」「答案」两列时，也可以直接导入。",
]


def build_xlsx_template(examples: List[Dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "知识条目"
    head_font = Font(bold=True, color="FFFFFF")
    req_fill = PatternFill("solid", fgColor="1E6FD9")
    opt_fill = PatternFill("solid", fgColor="5B7A99")
    widths = {"title": 22, "category": 12, "triggers": 26, "scenario": 26, "steps": 30,
              "principles": 24, "example_reply_zh": 40, "forbidden": 20, "reply_mode": 12,
              "negative_triggers": 16, "enabled": 8}
    for ci, (field, zh, req, hint) in enumerate(SHEET_COLUMNS, start=1):
        c = ws.cell(row=1, column=ci, value=zh)
        c.font = head_font
        c.fill = req_fill if req else opt_fill
        c.alignment = Alignment(vertical="center")
        c.comment = Comment(hint, "智聊")
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(field, 16)
    for ri, ex in enumerate(examples, start=2):
        for ci, v in enumerate(_template_row(ex), start=1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(color="7A7A7A", italic=True)
    ws.freeze_panes = "B2"
    col = {f: get_column_letter(i) for i, (f, _z, _r, _h) in enumerate(SHEET_COLUMNS, start=1)}
    dv_mode = DataValidation(type="list", formula1='"AI参考,AI严格,直接输出"', allow_blank=True)
    dv_mode.add(f"{col['reply_mode']}2:{col['reply_mode']}2000")
    dv_on = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    dv_on.add(f"{col['enabled']}2:{col['enabled']}2000")
    ws.add_data_validation(dv_mode)
    ws.add_data_validation(dv_on)

    guide = wb.create_sheet("填写说明")
    guide.append(["列名", "必填", "说明"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for _field, zh, req, hint in SHEET_COLUMNS:
        guide.append([zh, "是" if req else "", hint])
    guide.append([])
    guide.append(["注意事项"])
    guide.cell(row=guide.max_row, column=1).font = Font(bold=True)
    for i, note in enumerate(TEMPLATE_NOTES, start=1):
        guide.append([f"{i}.", "", note])
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 6
    guide.column_dimensions["C"].width = 80
    for row in guide.iter_rows(min_row=2):
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
