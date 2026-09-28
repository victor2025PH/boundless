# -*- coding: utf-8 -*-
"""知识库表格导入 + 导入模板门禁。

钉住的事故形态：
  - CSV 导出只选 8 列却写 12 列表头，更新模式导回 → reply_mode / template_key 等被清成空串；
  - CSV 导入把触发词转成 JSON 字符串再交给 add_entry 按逗号切 → 触发词变成 '["退款"'；
  - 无标题行被静默丢弃，结果里既不算失败也没有行号。
"""
from __future__ import annotations

import csv
import io
import json

import pytest

from src.utils import kb_sheet_io as sio
from src.utils.kb_store import KnowledgeBaseStore


@pytest.fixture()
def store(tmp_path):
    return KnowledgeBaseStore(tmp_path / "kb.db")


def _by_title(store, title):
    for e in store.list_entries():
        if e["title"] == title:
            e = dict(e)
            for f in ("triggers", "negative_triggers"):
                if isinstance(e.get(f), str):
                    e[f] = json.loads(e[f] or "[]")
            return e
    return None


def test_split_terms_accepts_all_separators_and_json():
    assert sio.split_terms("退款; 退货，换货、发票|开票") == ["退款", "退货", "换货", "发票", "开票"]
    assert sio.split_terms('["a", "b"]') == ["a", "b"]
    assert sio.split_terms(["x", " ", "y"]) == ["x", "y"]
    assert sio.split_terms(None) == [] and sio.split_terms("") == []


def test_add_entry_json_string_triggers_not_corrupted(store):
    store.add_entry({"title": "退款流程", "triggers": json.dumps(["退款", "如何退款"], ensure_ascii=False)})
    assert _by_title(store, "退款流程")["triggers"] == ["退款", "如何退款"]


def test_csv_roundtrip_update_keeps_advanced_fields(store, tmp_path):
    store.add_entry({"title": "直出条目", "triggers": ["价格"], "reply_mode": "direct",
                     "template_key": "price_tpl", "fallback_group": "g1",
                     "negative_triggers": ["免费"], "example_reply_zh": "原回复"})
    text = store.export_csv()
    rows = list(csv.DictReader(io.StringIO(text.lstrip("\ufeff"))))
    assert rows[0]["reply_mode"] == "direct" and rows[0]["template_key"] == "price_tpl"
    assert rows[0]["negative_triggers"] == "免费"

    r = store.import_from_csv(text, mode="update")
    assert r["updated"] == 1 and r["failed"] == 0
    e = _by_title(store, "直出条目")
    assert e["reply_mode"] == "direct"
    assert e["template_key"] == "price_tpl"
    assert e["fallback_group"] == "g1"
    assert e["triggers"] == ["价格"]
    assert e["negative_triggers"] == ["免费"]

    other = KnowledgeBaseStore(tmp_path / "kb2.db")
    r2 = other.import_from_csv(text)
    assert r2["added"] == 1
    assert _by_title(other, "直出条目")["triggers"] == ["价格"]


def test_update_with_blank_cells_does_not_clear(store):
    store.add_entry({"title": "已有", "scenario": "原场景", "steps": "原步骤", "enabled": 0})
    r = store.import_from_csv("标题,使用场景,处理步骤\n已有,新场景,\n", mode="update")
    assert r["updated"] == 1
    e = _by_title(store, "已有")
    assert e["scenario"] == "新场景" and e["steps"] == "原步骤"
    assert int(e["enabled"]) == 0


def test_chinese_headers_aliases_and_values(store):
    text = ("\ufeff标题（必填）,分类,触发词,标准回复,回复模式,启用,备注\n"
            "发货时间,常规咨询,多久发货；什么时候发货、发货,一般48小时内发货,直接输出,否,内部\n")
    r = store.import_from_csv(text)
    assert r["added"] == 1
    assert r["unknown_headers"] == ["备注"]
    e = _by_title(store, "发货时间")
    assert e["triggers"] == ["多久发货", "什么时候发货", "发货"]
    assert e["reply_mode"] == "direct"
    assert int(e["enabled"]) == 0
    assert e["source"] == "import"


def test_faq_question_answer_sheet_gets_triggers(store):
    r = store.import_from_csv("问题,答案\n怎么退款,在订单页点申请退款\n")
    assert r["added"] == 1
    e = _by_title(store, "怎么退款")
    assert e["triggers"] == ["怎么退款"]
    assert e["example_reply_zh"] == "在订单页点申请退款"


def test_row_level_errors_and_warnings_have_row_numbers(store):
    text = "标题,触发词,回复模式\n好条目,a,AI严格\n,b,\n\n坏模式,c,乱写\n无触发词,,\n"
    r = store.import_from_csv(text)
    assert r["added"] == 3
    assert r["failed"] == 1
    assert r["errors"] == [{"row": 3, "title": "", "reason": "no_title"}]
    reasons = {(w["row"], w["reason"]) for w in r["warnings"]}
    assert (5, "bad_reply_mode") in reasons and (6, "no_triggers") in reasons
    assert _by_title(store, "好条目")["reply_mode"] == "ai_strict"


def test_missing_title_column(store):
    r = store.import_from_csv("分类,触发词\nA,b\n")
    assert r["missing_title_column"] is True
    assert r["added"] == 0 and store.list_entries() == []


def test_dry_run_writes_nothing(store):
    store.add_entry({"title": "已有", "triggers": ["x"]})
    text = "标题,触发词\n已有,y\n新条目,z\n新条目,z2\n"
    r = store.import_from_csv(text, mode="update", dry_run=True)
    assert r["dry_run"] is True
    assert (r["added"], r["updated"], r["total"]) == (1, 2, 3)
    assert [p["action"] for p in r["preview"]] == ["update", "add", "update"]
    assert len(store.list_entries()) == 1
    assert _by_title(store, "已有")["triggers"] == ["x"]


def test_csv_template_examples_skipped_by_default(store):
    tpl = sio.build_csv_template(sio.template_examples("sales"))
    assert tpl.startswith("\ufeff标题,分类,触发词")
    r = store.import_from_csv(tpl)
    assert r["added"] == 0 and r["examples_skipped"] == 2
    r2 = store.import_from_csv(tpl, skip_examples=False)
    assert r2["added"] == 2


def test_xlsx_template_roundtrip(store):
    from openpyxl import load_workbook
    data = sio.build_xlsx_template(sio.template_examples("companion"))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["知识条目", "填写说明"]
    ws = wb["知识条目"]
    assert ws["A1"].value == "标题" and ws["A1"].comment is not None
    assert ws.max_row == 3
    r = store.import_from_xlsx(data, skip_examples=False)
    assert r["added"] == 2 and r["failed"] == 0
    assert all(e["triggers"] for e in (_by_title(store, t["title"]) for t in r["preview"]))


@pytest.fixture()
def client(store, tmp_path):
    from types import SimpleNamespace

    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from starlette.middleware.sessions import SessionMiddleware

    from src.web.routes.kb_routes import register_kb_routes

    app = FastAPI()
    app.add_middleware(SessionMiddleware, secret_key="t")
    cfg = SimpleNamespace(config_path=str(tmp_path / "config.yaml"), config={},
                          get=lambda key, default=None: default)
    ctx = SimpleNamespace(kb_store=store, config_manager=cfg, audit_store=None,
                          api_auth=lambda request: None, require_auth=lambda request: None,
                          fire_webhook=lambda *a, **k: None)
    register_kb_routes(app, ctx)
    return TestClient(app)


def test_route_add_entry_auto_embeds_when_endpoint_available(client, store, monkeypatch):
    """P1-3：保存即向量化——有嵌入端点（网关 /api/ai/v1）时新建 / 改匹配面后台重算这一条；
    没有端点（裸 DeepSeek 对话端点）一枪不打。"""
    calls = []
    # 路由闭包里的 _call_embed_api 不可直接替换 → 在 httpx 层打桩
    import httpx

    class _Resp:
        def json(self):
            return {"data": [{"index": 0, "embedding": [0.1, 0.2, 0.3]}]}

    class _Client:
        def __init__(self, *a, **k):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *a):
            return False

        async def post(self, url, **kw):
            calls.append((url, kw.get("json", {}).get("input")))
            return _Resp()

    monkeypatch.setattr(httpx, "AsyncClient", _Client)
    # 先无端点：不打
    r = client.post("/api/kb/entries", json={"title": "网站有什么活动", "triggers": ["活动"],
                                             "category": "活动优惠", "example_reply_zh": "新人礼"}).json()
    assert r["ok"] and calls == []
    # 配网关端点：打一枪，向量落库
    from types import SimpleNamespace
    cm = SimpleNamespace(config={"ai": {"base_url": "https://bd2026.cc/api/ai/v1", "api_key": "cx.x"}},
                         config_path="x")
    # 找到路由闭包里的 config_manager 引用：register_kb_routes 用 ctx.config_manager，
    # 这里直接改 ctx 对象不可达 → 用第二个 client 重新注册
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from starlette.middleware.sessions import SessionMiddleware
    from src.web.routes.kb_routes import register_kb_routes
    app2 = FastAPI()
    app2.add_middleware(SessionMiddleware, secret_key="t")
    ctx2 = SimpleNamespace(kb_store=store, config_manager=cm, audit_store=None,
                           api_auth=lambda request: None, require_auth=lambda request: None,
                           fire_webhook=lambda *a, **k: None)
    register_kb_routes(app2, ctx2)
    c2 = TestClient(app2)
    r2 = c2.post("/api/kb/entries", json={"title": "怎么注册", "triggers": ["注册"],
                                         "category": "注册登录", "example_reply_zh": "右上角注册"}).json()
    assert r2["ok"]
    # 后台任务在同一事件循环里排队；再发一个请求让循环跑完
    c2.get("/api/kb/embed-progress")
    assert calls and calls[-1][0].endswith("/embeddings") and "注册" in calls[-1][1][0]
    assert store.get_entry(r2["id"]).get("embedding")


def test_route_new_entry_templates_follow_persona_kinds(client, monkeypatch):
    """P1-2：新建模板端点按在用人设 kind 前置客服示例，分类表带客服集。"""
    import src.utils.persona_kind as pk
    monkeypatch.setattr(pk, "kinds_in_use", lambda cfg=None: ("support",))
    r = client.get("/api/kb/new-entry-templates").json()
    assert r["kinds"] == ["support"]
    assert r["templates"][0]["key"] == "support_promo"
    assert r["templates"][0]["category"] == "活动优惠"
    monkeypatch.setattr(pk, "kinds_in_use", lambda cfg=None: ())
    r2 = client.get("/api/kb/new-entry-templates").json()
    assert r2["templates"][0]["key"] != "support_promo"


def test_route_template_downloads(client):
    r = client.get("/api/kb/import-template?fmt=xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "spreadsheetml" in r.headers["content-type"]
    r = client.get("/api/kb/import-template?fmt=csv")
    assert r.content.startswith("\ufeff标题".encode("utf-8"))
    r = client.get("/api/kb/import-template?fmt=json")
    assert r.json()["entries"]


def test_route_xlsx_dry_run_then_import(client, store):
    import base64
    data = sio.build_xlsx_template(sio.template_examples("sales"))
    b64 = base64.b64encode(data).decode()
    r = client.post("/api/kb/import-csv", json={"xlsx_b64": b64, "dry_run": True,
                                                "skip_examples": False}).json()
    assert r["dry_run"] is True and r["added"] == 2 and store.list_entries() == []
    r = client.post("/api/kb/import-csv", json={"xlsx_b64": b64}).json()
    assert r["added"] == 0 and r["examples_skipped"] == 2
    r = client.post("/api/kb/import-csv", json={"csv": "标题,分类\n甲,不存在的分类\n"}).json()
    assert r["added"] == 1 and r["unknown_categories"] == ["不存在的分类"]


def test_route_yaml_text_import(client, store):
    text = "version: '1.0'\nentries:\n  - title: 年假\n    triggers: [年假, 休假]\n"
    r = client.post("/api/kb/import", json={"text": text, "format": "yaml", "dry_run": True}).json()
    assert r["added"] == 1 and store.list_entries() == []
    r = client.post("/api/kb/import", json={"text": text, "format": "yaml"}).json()
    assert r["added"] == 1
    assert _by_title(store, "年假")["triggers"] == ["年假", "休假"]
    bad = client.post("/api/kb/import", json={"text": "{not json", "format": "json"})
    assert bad.status_code == 400


def test_import_batch_recorded_and_undone(store):
    store.add_entry({"title": "已有", "triggers": ["旧"], "scenario": "原场景"})
    assert store.import_from_csv("标题,触发词\n已有,新\n", mode="update", dry_run=True).get("batch_id") is None
    r = store.import_from_csv("标题,触发词,使用场景\n已有,新,新场景\n新甲,a,\n新乙,b,\n",
                              mode="update", batch={"filename": "faq.csv", "fmt": "csv",
                                                    "operator": "admin"})
    assert (r["added"], r["updated"]) == (2, 1) and r["batch_id"]
    [b] = store.list_import_batches()
    assert (b["filename"], b["added"], b["updated"], b["undone_at"]) == ("faq.csv", 2, 1, "")

    u = store.undo_import_batch(r["batch_id"])
    assert (u["ok"], u["deleted"], u["restored"], u["kept"], u["remaining"]) == (True, 2, 1, [], 0)
    assert [e["title"] for e in store.list_entries()] == ["已有"]
    e = _by_title(store, "已有")
    assert e["triggers"] == ["旧"] and e["scenario"] == "原场景"
    assert store.list_import_batches()[0]["undone_at"]
    assert store.undo_import_batch(r["batch_id"]) == {"ok": False, "reason": "already_undone"}
    assert store.undo_import_batch("nope")["reason"] == "not_found"


def test_undo_keeps_entries_edited_after_import_unless_forced(store):
    r = store.import_from_csv("标题,触发词\n甲,a\n乙,b\n", batch={"filename": "x.csv"})
    edited = _by_title(store, "甲")
    store.update_entry(edited["id"], {"scenario": "导入后手工补的"})
    u = store.undo_import_batch(r["batch_id"])
    assert (u["deleted"], u["kept"], u["remaining"]) == (1, ["甲"], 1)
    assert [e["title"] for e in store.list_entries()] == ["甲"]
    assert not store.list_import_batches()[0]["undone_at"]
    u2 = store.undo_import_batch(r["batch_id"], force=True)
    assert (u2["deleted"], u2["kept"]) == (1, [])
    assert store.list_entries() == []


def test_skip_only_import_records_no_batch(store):
    store.add_entry({"title": "已有", "triggers": ["x"]})
    r = store.import_from_csv("标题,触发词\n已有,y\n")
    assert r["skipped"] == 1 and "batch_id" not in r
    assert store.list_import_batches() == []


def test_export_include_disabled_roundtrip(store, tmp_path):
    store.add_entry({"title": "启用的", "triggers": ["a"]})
    store.add_entry({"title": "停用的", "triggers": ["b"], "enabled": 0})
    assert [e["title"] for e in store.export_all()["entries"]] == ["启用的"]
    full = store.export_all(include_disabled=True)
    assert {e["title"] for e in full["entries"]} == {"启用的", "停用的"}
    text = store.export_csv(include_disabled=True)
    assert "enabled" in text.splitlines()[0] and "停用的" in text
    other = KnowledgeBaseStore(tmp_path / "kb3.db")
    assert other.import_from_csv(text)["added"] == 2
    assert int(_by_title(other, "停用的")["enabled"]) == 0
    other2 = KnowledgeBaseStore(tmp_path / "kb4.db")
    assert other2.import_from_data(full)["added"] == 2
    assert int(_by_title(other2, "停用的")["enabled"]) == 0


def test_route_import_history_and_undo(client, store):
    r = client.post("/api/kb/import-csv", json={"csv": "标题\n甲\n", "filename": "a.csv"}).json()
    hist = client.get("/api/kb/import-batches").json()["batches"]
    assert hist[0]["id"] == r["batch_id"] and hist[0]["filename"] == "a.csv"
    u = client.post(f"/api/kb/import-batches/{r['batch_id']}/undo", json={}).json()
    assert u["deleted"] == 1 and store.list_entries() == []
    again = client.post(f"/api/kb/import-batches/{r['batch_id']}/undo", json={})
    assert again.status_code == 409
    assert client.post("/api/kb/import-batches/zzz/undo", json={}).status_code == 404
    store.add_entry({"title": "停用的", "enabled": 0})
    assert "停用的" not in client.get("/api/kb/export-csv").content.decode("utf-8")
    assert "停用的" in client.get("/api/kb/export-csv?all=1").content.decode("utf-8")
    assert "（已停用）" in client.get("/api/kb/export-markdown?all=1").content.decode("utf-8")
    assert len(client.get("/api/kb/export?all=1").json()["entries"]) == 1


def test_import_ignores_file_ids_so_other_entries_are_not_overwritten(store):
    keep_id = store.add_entry({"title": "原条目", "steps": "别被覆盖"})
    r = store.import_from_data({"entries": [{"id": keep_id, "title": "备份里改过名的",
                                             "steps": "x"}]}, batch={"filename": "b.json"})
    assert r["added"] == 1
    assert store.get_entry(keep_id)["title"] == "原条目"
    store.undo_import_batch(r["batch_id"])
    assert store.get_entry(keep_id)["steps"] == "别被覆盖"
    assert [e["title"] for e in store.list_entries()] == ["原条目"]


def test_list_search_matches_triggers(store):
    store.add_entry({"title": "退款流程", "triggers": ["怎么退钱"]})
    store.add_entry({"title": "发货", "triggers": ["物流"]})
    assert [e["title"] for e in store.list_entries(search="退钱")] == ["退款流程"]
    assert [e["title"] for e in store.list_entries(search="物流")][0] == "发货"


def test_route_list_carries_translations_not_embeddings(client, store):
    a = store.add_entry({"title": "甲", "triggers": ["a"]})
    store.add_entry({"title": "乙", "triggers": ["b"]})
    store.upsert_translation(a, "en", {"title": "A"})
    with store._conn() as c:
        c.execute("UPDATE kb_entries SET embedding='[0.1,0.2]'")
    rows = {e["title"]: e for e in client.get("/api/kb/entries").json()["entries"]}
    assert set(rows["甲"]["translations"]) == {"en"}
    assert rows["乙"]["translations"] is None
    assert all("embedding" not in e for e in rows.values())


def test_json_template_importable(store):
    tpl = sio.build_json_template(sio.template_examples("sales"))
    assert tpl["version"] == "1.0" and len(tpl["entries"]) == 2
    assert store.import_from_data(tpl, skip_examples=True)["examples_skipped"] == 2
    assert store.import_from_data(tpl)["added"] == 2
