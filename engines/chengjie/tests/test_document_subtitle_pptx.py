"""P1（2026-08-18）：.srt/.vtt 字幕保时间轴翻译 + .pptx 保版式翻译单测。

风格对齐 test_document_file_translate.py：真 python-pptx 构造内存文档 + stub
TranslationService。字幕分类器是纯函数 → 金标逐行断言（时间轴/序号/cue id/
元块一个都不许被译；文本行一个都不许漏）。
"""
from io import BytesIO

import pytest

from src.ai.document_file_translate import (
    classify_subtitle_lines,
    translate_subtitle,
)
from src.ai.translation_engines import EngineResult, EngineRouter
from src.ai.translation_service import TranslationService


class _StubEngine:
    def __init__(self, name="ai", *, fail_on=None):
        self.name = name
        self._fail_on = set(fail_on or [])

    @property
    def available(self):
        return True

    def supports_target(self, t):
        return True

    async def translate(self, text, *, source_lang, target_lang, style="chat", glossary_hint=""):
        if text in self._fail_on:
            return EngineResult("", self.name, False, error="boom")
        return EngineResult(f"{text}#{self.name}", self.name, True)


def _svc(engine=None):
    s = TranslationService(ai_client=None)
    s._router = EngineRouter([engine or _StubEngine()])
    return s


_SRT = (
    "1\n"
    "00:00:01,000 --> 00:00:03,500\n"
    "Hello there\n"
    "\n"
    "2\n"
    "00:00:04,000 --> 00:00:06,000\n"
    "Second line one\n"
    "Second line two\n"
    "\n"
)

_VTT = (
    "WEBVTT - demo\n"
    "\n"
    "NOTE\n"
    "this block must be preserved\n"
    "\n"
    "intro-cue\n"
    "00:00:01.000 --> 00:00:03.000\n"
    "Hello vtt\n"
    "\n"
    "00:00:04.000 --> 00:00:06.000\n"
    "Plain cue text\n"
)


# ── 分类器金标 ─────────────────────────────────────────────────────────


def test_classify_srt_gold():
    lines = _SRT.split("\n")
    flags = classify_subtitle_lines(lines, kind="srt")
    translated = [ln for ln, f in zip(lines, flags) if f]
    assert translated == ["Hello there", "Second line one", "Second line two"]


def test_classify_vtt_gold():
    lines = _VTT.split("\n")
    flags = classify_subtitle_lines(lines, kind="vtt")
    translated = [ln for ln, f in zip(lines, flags) if f]
    # WEBVTT 头 / NOTE 块 / cue id（intro-cue，下一行是时间轴）/ 时间轴全部保留
    assert translated == ["Hello vtt", "Plain cue text"]


def test_classify_numeric_only_line_is_preserved():
    # 纯数字行即便不挨时间轴（罕见畸形），也按序号保留——宁可漏译不误改结构
    flags = classify_subtitle_lines(["42", "text line"], kind="srt")
    assert flags == [False, True]


# ── srt/vtt 端到端 ────────────────────────────────────────────────────


async def test_translate_srt_keeps_structure():
    res = await translate_subtitle(
        _SRT.encode("utf-8"), xlate=_svc(), kind="srt",
        target_lang="zh", source_lang="en")
    assert res["ok"] is True
    out = res["data"].decode("utf-8")
    assert "00:00:01,000 --> 00:00:03,500" in out      # 时间轴逐字保留
    assert "Hello there#ai" in out
    assert "Second line two#ai" in out
    assert out.split("\n")[0] == "1"                   # 序号行原位
    assert res["stats"]["translated"] == 3


async def test_translate_vtt_preserves_meta_blocks():
    res = await translate_subtitle(
        _VTT.encode("utf-8"), xlate=_svc(), kind="vtt",
        target_lang="zh", source_lang="en")
    assert res["ok"] is True
    out = res["data"].decode("utf-8")
    assert out.startswith("WEBVTT - demo")
    assert "this block must be preserved" in out       # NOTE 块未被译
    assert "intro-cue\n" in out                        # cue id 未被译
    assert "Hello vtt#ai" in out and "Plain cue text#ai" in out


async def test_translate_srt_bilingual_stacks_lines():
    res = await translate_subtitle(
        _SRT.encode("utf-8"), xlate=_svc(), kind="srt",
        target_lang="zh", source_lang="en", bilingual=True)
    out = res["data"].decode("utf-8")
    assert "Hello there\nHello there#ai" in out        # 原文行在上、译文行在下
    assert "00:00:01,000 --> 00:00:03,500" in out


async def test_translate_srt_failure_keeps_original_line():
    res = await translate_subtitle(
        _SRT.encode("utf-8"), xlate=_svc(_StubEngine(fail_on=["Hello there"])),
        kind="srt", target_lang="zh", source_lang="en")
    assert res["ok"] is True
    out = res["data"].decode("utf-8")
    assert "Hello there" in out and "Hello there#ai" not in out
    assert res["stats"]["failed"] == 1


async def test_translate_srt_bom_and_crlf_tolerated():
    raw = ("\ufeff" + _SRT.replace("\n", "\r\n")).encode("utf-8")
    res = await translate_subtitle(raw, xlate=_svc(), kind="srt",
                                   target_lang="zh", source_lang="en")
    assert res["ok"] is True
    assert "Hello there#ai" in res["data"].decode("utf-8")


async def test_translate_subtitle_no_text_soft_fail():
    only_meta = "1\n00:00:01,000 --> 00:00:02,000\n\n"
    res = await translate_subtitle(only_meta.encode("utf-8"), xlate=_svc(),
                                   kind="srt", target_lang="zh")
    assert res["ok"] is False and res["reason"] == "no_text"


# ── pptx ──────────────────────────────────────────────────────────────
pptx = pytest.importorskip("pptx")

from src.ai.document_file_translate import translate_pptx  # noqa: E402


def _make_pptx(texts, *, with_notes=""):
    from pptx.util import Inches
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # 空白版式
    for i, t in enumerate(texts):
        box = slide.shapes.add_textbox(Inches(1), Inches(0.5 + i), Inches(4), Inches(0.8))
        box.text_frame.text = t
    if with_notes:
        slide.notes_slide.notes_text_frame.text = with_notes
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pptx_texts(data):
    prs = pptx.Presentation(BytesIO(data))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    out.append(p.text)
        if slide.has_notes_slide:
            out.append(slide.notes_slide.notes_text_frame.text)
    return out


async def test_translate_pptx_textboxes_and_notes():
    data = _make_pptx(["Hello", "World"], with_notes="Speaker note")
    res = await translate_pptx(data, xlate=_svc(), target_lang="zh", source_lang="en")
    assert res["ok"] is True
    texts = _pptx_texts(res["data"])
    assert "Hello#ai" in texts and "World#ai" in texts
    assert any("Speaker note#ai" in t for t in texts)


async def test_translate_pptx_failure_keeps_original():
    data = _make_pptx(["keep", "bad"])
    res = await translate_pptx(data, xlate=_svc(_StubEngine(fail_on=["bad"])),
                               target_lang="zh", source_lang="en")
    assert res["ok"] is True
    texts = _pptx_texts(res["data"])
    assert "keep#ai" in texts and "bad" in texts
    assert res["stats"]["failed"] == 1


async def test_bad_pptx_soft_fail():
    res = await translate_pptx(b"not a pptx", xlate=_svc(), target_lang="zh")
    assert res["ok"] is False and res["reason"] == "bad_pptx"


async def test_pptx_reopenable_output():
    data = _make_pptx(["Reopen me"])
    res = await translate_pptx(data, xlate=_svc(), target_lang="zh", source_lang="en")
    reopened = pptx.Presentation(BytesIO(res["data"]))
    assert reopened.slides
